#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把各厂区 Outlook 值班/巡查 .msg 汇总成与「汇总.xlsx」相同版式的 Excel。

用法:
    python summarize_duty_reports.py --input 邮件目录 --start 2026-07-09 --end 2026-07-09 --output 汇总结果.xlsx

可选核对（手里有人工示例时）:
    python summarize_duty_reports.py --input 邮件目录 --start 2026-07-09 --end 2026-07-09 \\
        --output 汇总结果.xlsx --check 汇总.xlsx

依赖: extract-msg、beautifulsoup4、openpyxl（见同目录 requirements.txt）
"""

from __future__ import annotations

import argparse
import re
import sys
import traceback
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path

try:
    from zoneinfo import ZoneInfo

    SHANGHAI = ZoneInfo("Asia/Shanghai")
except Exception:  # pragma: no cover
    SHANGHAI = timezone(timedelta(hours=8))

COLUMNS = [
    "来源",
    "发件人",
    "巡查类型",
    "日期",
    "厂区",
    "属地",
    "责任区域",
    "第一责任人",
    "责任人",
    "隐患分级",
    "事件描述",
    "隐患类型",
]

# 表头单元格去掉空白后的精确匹配。长名称放前面，避免「责任人」吃掉「第一责任人」。
HEADER_FIELDS = [
    ("第一责任人", {"第一责任人", "属地负责人"}),
    ("责任区域", {"责任区域"}),
    ("责任人", {"责任人"}),
    ("巡查类型", {"巡查类别", "巡查项类别", "巡查类型"}),
    ("隐患分级", {"隐患分级"}),
    ("事件描述", {"巡查结果", "巡查内容/状态", "巡查内容和状态", "巡查发现项", "发现项"}),
    ("属地", {"属地", "建筑物"}),
    ("厂区", {"厂区"}),
    ("日期", {"日期", "值班日期"}),
    ("序号", {"序号"}),
    ("内容", {"内容"}),
    ("关注内容", {"关注内容"}),
    ("巡查位置", {"巡查位置"}),
]

COL_WIDTHS = {
    "A": 15.625,
    "B": 14.75,
    "C": 10.625,
    "D": 13.0,
    "E": 6.625,
    "F": 13.0,
    "G": 13.0,
    "H": 13.0,
    "I": 13.0,
    "J": 16.0,
    "K": 55.125,
    "L": 17.125,
}

PLANT_RE = re.compile(r"(DH\d+|TJ\d+|FX\d+|SH\d+)", re.I)
GRADE_OK = {"一般隐患", "重大隐患"}

# 隐患类型：按顺序命中即停。只用于最后一列，允许不准。
HAZARD_RULES = [
    ("静电", ("静电",)),
    ("氮气保护", ("氮气",)),
    ("气味扩散", ("气味", "异味")),
    ("标识标牌", ("标识", "标签", "标牌", "色标")),
    ("文件记录", ("记录", "合格签", "有效期", "过期", "审批")),
    ("跑冒滴漏", ("滴漏", "渗漏", "漏水", "漏油", "漏气", "渗液", "析出", "滴水", "滴落", "溢水", "渗漏", "漏")),
    ("跑冒滴漏（预防）", ("单阀", "未上锁", "未关闭", "盲板", "未密封", "缺少盲板")),
    ("行为/操作", ("刷视频", "抽烟", "未佩戴", "没有佩戴", "脱岗", "操作习惯", "劳保", "返岗")),
    ("高风险操作核查", ("checklist", "投料", "掏", "减压", "浓缩", "敞口", "直接排放", "冲釜", "压滤")),
    ("设备异常", ("故障", "损坏", "报警", "松脱", "断裂", "失灵", "失效", "高报")),
    ("现场5S", ("清理", "粉尘", "锈", "卫生", "杂乱", "摆放", "积油", "印记")),
    ("物料状态", ("暂存", "入库", "物料暴露")),
    ("桶/槽子操作", ("塑料桶", "方槽", "吨桶", "吨罐")),
    ("设备设施隐患识别", ("缺少", "未安装", "变形", "脱落", "开口", "不完善", "管路", "阀门")),
    ("现场状态", ("现场状态", "在岗")),
]


@dataclass
class MailMeta:
    path: Path
    subject: str
    sender_raw: str
    sender: str
    sent_at: datetime | None
    body: str
    html: str
    filename_dates: list[date] = field(default_factory=list)
    subject_dates: list[date] = field(default_factory=list)
    header_dates: list[date] = field(default_factory=list)
    table_dates: list[date] = field(default_factory=list)
    is_reply: bool = False
    base_subject: str = ""

    @property
    def content_dates(self) -> list[date]:
        return self.filename_dates + self.subject_dates + self.header_dates + self.table_dates

    @property
    def sent_date(self) -> date | None:
        if not self.sent_at:
            return None
        return self.sent_at.astimezone(SHANGHAI).date()

    @property
    def sent_hour(self) -> int | None:
        if not self.sent_at:
            return None
        return self.sent_at.astimezone(SHANGHAI).hour


@dataclass
class Row:
    来源: str = ""
    发件人: str = ""
    巡查类型: str = ""
    日期: date | None = None
    厂区: str = ""
    属地: str = ""
    责任区域: str = ""
    第一责任人: str = ""
    责任人: str = ""
    隐患分级: str = ""
    事件描述: str = ""
    隐患类型: str = ""
    source_file: str = ""

    def as_dict(self) -> dict:
        return {k: getattr(self, k) for k in COLUMNS}


def clean_text(value: str | None) -> str:
    if not value:
        return ""
    text = str(value).replace("\xa0", " ").replace("\u3000", " ").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n+", " ", text)
    return text.strip()


def tighten(value: str | None) -> str:
    """去掉 Word 把相邻文字拆开后多出来的空格，保留标点后的单个空格。"""
    text = clean_text(value)
    if not text:
        return ""
    text = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", text)
    text = re.sub(r"(?<=[A-Za-z0-9])\s+(?=[\u4e00-\u9fff])", "", text)
    text = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[A-Za-z0-9])", "", text)
    text = re.sub(r"(?<=[A-Za-z])\s+(?=[A-Za-z])", "", text)
    text = re.sub(r"\s*([/&])\s*", r"\1", text)
    text = re.sub(r"_\s+", "_", text)
    text = re.sub(r"\s*-\s*", "-", text)
    text = re.sub(r"\s+([，。；、：,.;）)])", r"\1", text)
    text = re.sub(r"([（(])\s+", r"\1", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip(" \t")


def norm_key(value: str | None) -> str:
    text = tighten(value).replace(" ", "")
    return text.replace("／", "/").replace("／", "/")


def fingerprint(value: str | None) -> str:
    text = tighten(value).lower()
    text = re.sub(r"\s+", "", text)
    text = text.replace("，", "").replace(",", "").replace("。", "").replace("；", "")
    text = text.replace(";", "").replace("：", "").replace(":", "").replace("、", "")
    text = text.replace("（", "(").replace("）", ")").replace("　", "")
    return text


def parse_dates(text: str | None) -> list[date]:
    if not text:
        return []
    found: list[date] = []
    seen: set[date] = set()

    def add(y: int, m: int, d: int) -> None:
        try:
            value = date(y, m, d)
        except ValueError:
            return
        if value.year < 2020 or value.year > 2100:
            return
        if value not in seen:
            seen.add(value)
            found.append(value)

    for match in re.finditer(r"(20\d{2})\s*[./\-年]\s*(\d{1,2})\s*[./\-月]\s*(\d{1,2})", text):
        add(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    for match in re.finditer(r"(?<!\d)(\d{1,2})\s*[./\- ]\s*(\d{1,2})\s*[./\- ]\s*(20\d{2})", text):
        add(int(match.group(3)), int(match.group(1)), int(match.group(2)))
    for match in re.finditer(r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)", text):
        add(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    return found


def sheet_title(start: date | None, end: date | None) -> str:
    def fmt(value: date) -> str:
        return f"{value.year}.{value.month}.{value.day}"

    if start and end:
        if start == end:
            return fmt(start)[:31]
        return f"{fmt(start)}-{fmt(end)}"[:31]
    if start:
        return fmt(start)[:31]
    return "值班汇总"


def decode_maybe(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bytes):
        for enc in ("utf-8", "gb18030", "gbk", "utf-16"):
            try:
                return value.decode(enc)
            except Exception:
                continue
        return value.decode("utf-8", "replace")
    return str(value)


def is_reply_subject(subject: str) -> bool:
    return bool(re.match(r"^\s*(答复|回复|re|fw|fwd|转发)\s*[:：]", subject or "", flags=re.I))


def base_subject(subject: str) -> str:
    text = subject or ""
    while True:
        updated = re.sub(r"^\s*(答复|回复|re|fw|fwd|转发)\s*[:：]\s*", "", text, flags=re.I)
        if updated == text:
            return text.strip()
        text = updated


NOT_PERSON = re.compile(r"主管|主任|经理|工程师|部门|生产|设备|工艺|凯莱|集团|公司|值班|序号|巡查|班|楼|厂|房")


def source_name(subject: str) -> str:
    text = base_subject(subject)
    text = text.replace("【值班报告】", "")
    text = re.sub(r"(20\d{2}\s*[./\-年]\s*\d{1,2}\s*[./\-月]\s*\d{1,2}日?)", "", text)
    text = re.sub(r"(?<!\d)(\d{1,2}\s*[./\-]\s*\d{1,2}\s*[./\-]\s*20\d{2})", "", text)
    text = re.sub(r"(?<!\d)(20\d{2}\d{2}\d{2})(?!\d)", "", text)

    def drop_person(match: re.Match) -> str:
        token = match.group(1)
        if NOT_PERSON.search(token):
            return match.group(0)
        return ""

    text = re.sub(r"[-—–]\s*([\u4e00-\u9fff]{2,4})(?=\s*[-—–]|$)", drop_person, text)
    text = re.sub(r"\s{2,}([\u4e00-\u9fff]{2,4})(?=\s*[-—–]|$)", drop_person, text)
    text = re.sub(r"(?:\s*[-—–]\s*){2,}", "", text)
    text = re.sub(r"\s+", " ", text).strip(" -—–_")
    text = re.sub(r"\s*[-—–]\s*$", "", text).strip()
    return text


def person_token(chunk: str | None) -> str:
    """从一小段文字里取出 2–4 个字的人名，停在「值班日期 / 序号」这类后接词前面。"""
    if not chunk:
        return ""
    chunk = re.split(r"值班日期|序号|巡查类别|一、|二、", chunk)[0]
    compact = re.sub(r"\s+", "", chunk)
    match = re.match(r"([\u4e00-\u9fff]{2,3})(?=序号|值班|巡查|$)|([\u4e00-\u9fff]{2,4})", compact)
    if not match:
        return ""
    name = match.group(1) or match.group(2)
    if NOT_PERSON.search(name):
        return ""
    return name


def chinese_person(text: str | None) -> str:
    name = person_token(text) or ""
    if name:
        return name
    fallback = re.sub(r"\s+", "", tighten(text))
    fallback = fallback.strip("：:/-—–| ")
    if re.fullmatch(r"[\u4e00-\u9fffA-Za-z0-9\-]{2,20}", fallback) and re.search(r"[\u4e00-\u9fff]", fallback):
        if not NOT_PERSON.search(fallback):
            return fallback
    return ""


def sender_name(subject: str, body: str, sender_raw: str) -> str:
    head = body[:2500]
    patterns = [
        r"值班人\s*员?\s*[:：]\s*([^\n]{1,120})",
        r"驻厂人员\s*/?\s*日期\s*[:：]\s*([^\n]{1,80})",
        r"值班人\s*[:：]\s*([^\n]{1,120})",
    ]
    for pattern in patterns:
        match = re.search(pattern, head)
        if match:
            name = person_token(match.group(1))
            if name:
                return name
    clipped = body.split("声明")[0].split("DISCLAIMER")[0]
    found = []
    for match in re.finditer(
        r"(?:^|\n)\s*([\u4e00-\u9fff](?:[ \t]*[\u4e00-\u9fff]){1,3})\s*[|/]\s*[A-Za-z]",
        clipped,
    ):
        name = person_token(match.group(1))
        if name:
            found.append(name)
    if found:
        return found[-1]
    match = re.search(r"[-—–]\s*([\u4e00-\u9fff]{2,4})\s*[-—–]", base_subject(subject))
    if match:
        name = person_token(match.group(1))
        if name:
            return name
    display = re.sub(r"<[^>]+>", "", sender_raw or "").strip().strip('"')
    if re.search(r"[\u4e00-\u9fff]", display):
        return person_token(display) or display
    return display


def header_dates(body: str) -> list[date]:
    head = body[:1800]
    blob = head[:400]
    for match in re.finditer(r".{0,12}(?:值班日期|驻厂人员|巡查日期).{0,30}", head):
        blob += "\n" + match.group(0)
    # 开头那句「以下为……07/09/2026值班报告」
    intro = re.search(r"以下为.{0,80}", head)
    if intro:
        blob += "\n" + intro.group(0)
    span = re.search(r"(\d{1,2}\s*/\s*\d{1,2}\s*[~～至到\-]\s*\d{1,2}\s*/\s*\d{1,2})", head)
    if span:
        blob += "\n" + span.group(0)
        # 补上年份：优先用文中出现的年份
        years = re.findall(r"20\d{2}", head[:500])
        year = years[0] if years else str(date.today().year)
        parts = re.findall(r"(\d{1,2})\s*/\s*(\d{1,2})", span.group(0))
        extra = []
        for month, day in parts:
            extra.extend(parse_dates(f"{year}-{int(month):02d}-{int(day):02d}"))
        return parse_dates(blob) + [d for d in extra if d not in parse_dates(blob)]
    return parse_dates(blob)


def plant_code(*texts: str) -> str:
    for text in texts:
        tightened = tighten(text).upper().replace(" ", "")
        if "敦化1" in text or "敦化一" in text:
            return "DH1"
        if "敦化2" in text or "敦化二" in text:
            return "DH2"
        if "敦化3" in text or "敦化三" in text:
            return "DH3"
        match = PLANT_RE.search(tightened)
        if match:
            return match.group(1).upper()
    return ""


def normalize_type(raw: str, subject: str, body: str = "") -> str:
    head = (body or "")[:1200]
    labeled = ""
    match = re.search(r"巡查类别\s*[:：]?\s*(外围值班|生产值班|高管驻场|高管驻厂|高管巡查)", head)
    if match:
        labeled = match.group(1)
    raw = raw or labeled
    text = f"{raw} {subject}"
    if "外围" in raw or "外围" in subject:
        return "外围值班"
    if any(token in raw for token in ("高管", "驻厂", "驻场")) or "驻厂" in subject or "驻场" in subject:
        if "生产值班" in raw and "驻" not in raw:
            return "生产值班"
        return "高管驻场"
    if "生产" in raw or "四点班" in subject or "生产值班" in subject or "四点班" in head:
        return "生产值班"
    if "值班" in subject and "驻厂" not in subject and "驻场" not in subject:
        return "生产值班"
    return tighten(raw)


def normalize_grade(raw: str) -> str:
    text = tighten(raw).replace(" ", "")
    if text in GRADE_OK:
        return text
    if "重大" in text:
        return "重大隐患"
    if "一般" in text:
        return "一般隐患"
    return ""


def classify_hazard(text: str) -> str:
    sample = tighten(text)
    if not sample:
        return ""
    lowered = sample.lower()
    for label, keys in HAZARD_RULES:
        for key in keys:
            if key.lower() in lowered:
                # 「无异常」里的「异」不要算成气味；「没有异常报警」里的报警仍算设备异常。
                if key in {"气味", "异味"} and "无异常" in sample and "气味" not in sample and "异味" not in sample:
                    continue
                if key == "漏" and not re.search(r"(漏水|漏油|漏气|渗漏|滴漏|泄漏|漏液|撒漏|渗液)", sample):
                    if "漏" not in sample:
                        continue
                    if not re.search(r"漏", sample):
                        continue
                return label
    return "其他"


def problem_re() -> re.Pattern:
    return re.compile(
        r"漏|渗|滴|坏|故障|缺|损|险|隐患|违规|脱落|报警|过期|堵塞|溢|敞|开口|未关|未启|气味|异味|锈|粉尘|变形|断裂|松脱|抽烟|刷视频|整改|不完善|盲板|高报"
    )


def has_problem(text: str) -> bool:
    return problem_re().search(tighten(text) or "") is not None


def is_headcount(text: str) -> bool:
    sample = tighten(text).strip("。；;，, ")
    if not sample or len(sample) > 28:
        return False
    if "在岗" not in sample:
        return False
    return re.search(r"漏|渗|坏|故障|缺|险|隐患|违规|损坏|脱落|报警|高报", sample) is None


def is_no_issue(text: str) -> bool:
    sample = tighten(text).strip("。；;，, ")
    if not sample:
        return True
    upper = sample.upper()
    if upper in {"N/A", "NA", "无", "无异常", "运行无异常", "巡查无异常"}:
        return True
    if re.fullmatch(r"(巡查期间)?无异常[。.]?(因.+变更)?", sample):
        return True
    problem = problem_re()
    no_issue_words = (
        "没有异常报警",
        "反馈设备没有异常",
        "反馈设备无异常",
        "反馈没有异常",
        "巡查无异常",
        "运行无异常",
        "没有异常",
        "未见异常",
        "未见违规",
        "无异常",
        "没有反应",
    )
    if any(word in sample for word in no_issue_words):
        remain = sample
        for word in no_issue_words:
            remain = remain.replace(word, "")
        if not problem.search(remain):
            return True
    if re.search(r"(项目无异常|人员工作无异常|人员无异常)", sample) and not problem.search(sample.replace("无异常", "")):
        return True
    return False


def split_numbered(text: str) -> list[tuple[str, str]]:
    sample = tighten(text)
    matches = list(re.finditer(r"(?<!\d)(\d{1,2})\s*[\.、．]\s*", sample))
    if not matches:
        return [("u", sample)] if sample else []
    chunks: list[tuple[str, str]] = []
    preamble = sample[: matches[0].start()].strip(" ；;。")
    if preamble:
        chunks.append(("u", preamble))
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(sample)
        body = sample[match.start() : end].strip(" ；;")
        if body:
            chunks.append(("n", body))
    if chunks and chunks[-1][0] == "n":
        kind, body = chunks[-1]
        pieces = [part.strip() for part in re.split(r"(?<=。)", body) if part.strip()]
        if len(pieces) >= 2:
            head = "".join(pieces[:-1]).strip()
            tail = pieces[-1].strip()
            # 编号条目后面又跟了一句没有编号的新问题，拆开。
            if head and len(tail) >= 8 and not re.match(r"^\d", tail) and re.search(r"坏|故障|缺|漏|脱落|报警|损坏|未", tail):
                chunks[-1] = (kind, head)
                chunks.append(("u", tail))
    return chunks


def split_buildings(text: str) -> list[tuple[str, str]]:
    sample = tighten(text)
    marks = list(re.finditer(r"(?<!\d)(\d{1,2}号楼)\s*[:：]", sample))
    if not marks:
        return [("", sample)] if sample else []
    blocks: list[tuple[str, str]] = []
    if marks[0].start() > 0:
        lead = sample[: marks[0].start()].strip()
        if lead:
            blocks.append(("", lead))
    for index, match in enumerate(marks):
        end = marks[index + 1].start() if index + 1 < len(marks) else len(sample)
        blocks.append((match.group(1), sample[match.end() : end].strip()))
    return blocks


def maybe_split_incidents(text: str) -> list[str]:
    parts = [part.strip("；; ") for part in re.split(r"[；;]", tighten(text)) if len(part.strip()) > 6]
    if len(parts) < 2:
        return [tighten(text)] if tighten(text) else []
    if not all(("——" in part or re.match(r"\d+名", part) or "抽烟" in part or "违规" in part) for part in parts):
        return [tighten(text)]
    output = []
    for index, part in enumerate(parts):
        if index == 0 and "——" in part and "：" in part:
            prefix, rest = part.split("：", 1)
            if 1 <= len(prefix) <= 8 and rest:
                part = rest
        output.append(part.strip())
    return output


def closed_change(text: str) -> bool:
    return re.search(r"无影响|已经事前提交变更|已经通过变更", tighten(text) or "") is not None


def keep_piece(kind: str, text: str) -> bool:
    sample = tighten(text).strip("。；; ")
    if len(sample) < 2 or is_headcount(sample):
        return False
    if kind == "n":
        return sample.upper() not in {"N/A", "NA", "无", "无异常"}
    return not is_no_issue(sample)


def load_messages(folder: Path) -> list[MailMeta]:
    import extract_msg

    metas: list[MailMeta] = []
    for path in sorted(folder.glob("*.msg")):
        message = extract_msg.Message(str(path))
        try:
            subject = decode_maybe(message.subject).strip()
            sender_raw = decode_maybe(message.sender).strip()
            body = decode_maybe(message.body)
            html = decode_maybe(message.htmlBody)
            sent = message.date
            if isinstance(sent, datetime) and sent.tzinfo is None:
                sent = sent.replace(tzinfo=timezone.utc)
            meta = MailMeta(
                path=path,
                subject=subject,
                sender_raw=sender_raw,
                sender=sender_name(subject, body, sender_raw),
                sent_at=sent if isinstance(sent, datetime) else None,
                body=body,
                html=html,
                filename_dates=parse_dates(path.name),
                subject_dates=parse_dates(subject),
                header_dates=header_dates(body),
                is_reply=is_reply_subject(subject),
                base_subject=base_subject(subject),
            )
            if html:
                # 只收集「日期」这一列，避免把事件描述里的过期日算进筛选范围。
                meta.table_dates = extract_table_dates(html)
            metas.append(meta)
        finally:
            message.close()
    return metas


def soup_of(html: str):
    from bs4 import BeautifulSoup

    return BeautifulSoup(html or "", "html.parser")


def direct_rows(table) -> list:
    rows = []
    for child in table.children:
        name = getattr(child, "name", None)
        if name == "tr":
            rows.append(child)
        elif name in {"tbody", "thead", "tfoot"}:
            rows.extend(child.find_all("tr", recursive=False))
    if not rows:
        rows = table.find_all("tr", recursive=False)
    return rows


def table_grid(table) -> list[list[str | None]]:
    rows = direct_rows(table)
    origin: dict[tuple[int, int], str] = {}
    covered: set[tuple[int, int]] = set()
    for r_index, tr in enumerate(rows):
        c_index = 0
        for cell in tr.find_all(["td", "th"], recursive=False):
            while (r_index, c_index) in covered or (r_index, c_index) in origin:
                c_index += 1
            rowspan = int(re.sub(r"\D", "", str(cell.get("rowspan") or "1")) or "1")
            colspan = int(re.sub(r"\D", "", str(cell.get("colspan") or "1")) or "1")
            text = clean_text(cell.get_text(" ", strip=True))
            origin[(r_index, c_index)] = text
            for dr in range(rowspan):
                for dc in range(colspan):
                    covered.add((r_index + dr, c_index + dc))
            c_index += colspan
    width = 0
    if covered:
        width = max(c for _, c in covered) + 1
    grid: list[list[str | None]] = []
    for r_index in range(len(rows)):
        line: list[str | None] = []
        for c_index in range(width):
            if (r_index, c_index) in origin:
                line.append(origin[(r_index, c_index)])
            elif (r_index, c_index) in covered:
                line.append(None)
            else:
                line.append("")
        grid.append(line)
    return grid


def header_map(cells: list[str | None]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, raw in enumerate(cells):
        if raw is None:
            continue
        key = norm_key(raw)
        if not key:
            continue
        for field_name, aliases in HEADER_FIELDS:
            if key in aliases and field_name not in mapping:
                mapping[field_name] = index
                break
    return mapping


def usable_header(mapping: dict[str, int]) -> bool:
    if "事件描述" not in mapping:
        return False
    return "属地" in mapping or "序号" in mapping


def find_header(grid: list[list[str | None]]) -> tuple[int, dict[str, int]] | None:
    best: tuple[int, dict[str, int]] | None = None
    for index, row in enumerate(grid):
        mapping = header_map(row)
        if not usable_header(mapping):
            continue
        score = len(mapping)
        if best is None or score >= len(best[1]):
            best = (index, mapping)
    return best


def cell_value(grid: list[list[str | None]], r_index: int, c_index: int | None, carry: bool) -> str:
    if c_index is None or r_index >= len(grid) or c_index >= len(grid[r_index]):
        return ""
    value = grid[r_index][c_index]
    if value is None and carry:
        for prev in range(r_index - 1, -1, -1):
            previous = grid[prev][c_index]
            if previous is not None:
                return previous
        return ""
    if value is None:
        return ""
    return value


def extract_table_dates(html: str) -> list[date]:
    soup = soup_of(html)
    dates: list[date] = []
    for table in soup.find_all("table"):
        if table.find("table"):
            continue
        grid = table_grid(table)
        found = find_header(grid)
        if not found:
            continue
        header_index, mapping = found
        if "日期" not in mapping:
            continue
        for r_index in range(header_index + 1, len(grid)):
            raw = cell_value(grid, r_index, mapping["日期"], carry=False)
            dates.extend(parse_dates(raw))
    return dates


def duty_anchor(meta: MailMeta) -> date | None:
    for pool in (meta.header_dates, meta.table_dates, meta.subject_dates, meta.filename_dates):
        if pool:
            return pool[0]
    return meta.sent_date


def parse_mail_rows(meta: MailMeta) -> list[Row]:
    if not meta.html:
        return []
    soup = soup_of(meta.html)
    rows: list[Row] = []
    for table in soup.find_all("table"):
        inner_hit = False
        for inner in table.find_all("table"):
            if find_header(table_grid(inner)):
                inner_hit = True
                break
        if inner_hit:
            continue
        grid = table_grid(table)
        found = find_header(grid)
        if not found:
            continue
        header_index, mapping = found
        header_cells = grid[header_index]
        desc_header = norm_key(header_cells[mapping["事件描述"]] if mapping["事件描述"] < len(header_cells) else "")
        checklist = desc_header == "巡查发现项" or "巡查位置" in mapping or "关注内容" in mapping
        grouped = "内容" in mapping
        if checklist:
            rows.extend(rows_from_checklist(meta, grid, header_index, mapping))
        elif grouped:
            rows.extend(rows_from_grouped(meta, grid, header_index, mapping))
        else:
            rows.extend(rows_from_flat(meta, grid, header_index, mapping))
    return dedupe_rows(rows)


def base_row(meta: MailMeta) -> Row:
    return Row(
        来源=source_name(meta.subject),
        发件人=meta.sender,
        厂区=plant_code(meta.subject, meta.body[:500]),
        source_file=meta.path.name,
    )


def apply_common(row: Row, meta: MailMeta, raw_type: str, raw_grade: str, raw_plant: str, raw_date: str) -> None:
    row.巡查类型 = normalize_type(raw_type, meta.subject, meta.body)
    row.隐患分级 = normalize_grade(raw_grade)
    plant = plant_code(raw_plant, row.属地, meta.subject)
    if plant:
        row.厂区 = plant
    parsed = parse_dates(raw_date)
    if parsed:
        row.日期 = parsed[0]
    row.属地 = tighten(row.属地)
    row.责任区域 = tighten(row.责任区域)
    if "/" in row.责任区域 and row.属地 and row.责任区域.split("/")[0] == row.属地:
        row.责任区域 = row.属地
    row.第一责任人 = chinese_person(row.第一责任人) or tighten(row.第一责任人)
    row.责任人 = tighten(row.责任人)
    if row.责任人.upper() in {"N/A", "NA"}:
        row.责任人 = ""
    if row.第一责任人.upper() in {"N/A", "NA"}:
        row.第一责任人 = ""
    row.隐患类型 = classify_hazard(row.事件描述)


def rows_from_flat(meta: MailMeta, grid, header_index: int, mapping: dict[str, int]) -> list[Row]:
    output: list[Row] = []
    for r_index in range(header_index + 1, len(grid)):
        desc_raw = cell_value(grid, r_index, mapping.get("事件描述"), carry=False)
        if not tighten(desc_raw):
            continue
        if norm_key(desc_raw) in {"巡查结果", "巡查内容/状态", "发现项", "巡查内容和状态"}:
            continue
        raw_grade = cell_value(grid, r_index, mapping.get("隐患分级"), carry=True)
        grade_blank = norm_key(raw_grade).upper() in {"", "N/A", "NA", "无"}
        pieces = []
        for piece in maybe_split_incidents(desc_raw):
            if is_no_issue(piece) or is_headcount(piece):
                continue
            # 分级写成 N/A 的行，只保留真有问题的句子，纯在岗/无异常说明不进汇总。
            if grade_blank and raw_grade and not has_problem(piece):
                continue
            pieces.append(piece)
        if not pieces:
            continue
        for piece in pieces:
            row = base_row(meta)
            row.属地 = cell_value(grid, r_index, mapping.get("属地"), carry=True)
            row.责任区域 = cell_value(grid, r_index, mapping.get("责任区域"), carry=True)
            row.第一责任人 = cell_value(grid, r_index, mapping.get("第一责任人"), carry=True)
            row.责任人 = cell_value(grid, r_index, mapping.get("责任人"), carry=True)
            row.事件描述 = tighten(piece)
            apply_common(
                row,
                meta,
                cell_value(grid, r_index, mapping.get("巡查类型"), carry=True),
                cell_value(grid, r_index, mapping.get("隐患分级"), carry=True),
                cell_value(grid, r_index, mapping.get("厂区"), carry=True),
                cell_value(grid, r_index, mapping.get("日期"), carry=True),
            )
            output.append(row)
    return output


def rows_from_grouped(meta: MailMeta, grid, header_index: int, mapping: dict[str, int]) -> list[Row]:
    output: list[Row] = []
    for r_index in range(header_index + 1, len(grid)):
        content = tighten(cell_value(grid, r_index, mapping.get("内容"), carry=False))
        desc_raw = cell_value(grid, r_index, mapping.get("事件描述"), carry=False)
        if "发现问题" not in content and "隐患" not in content:
            continue
        if not tighten(desc_raw) or is_no_issue(desc_raw):
            continue
        for kind, piece in split_numbered(desc_raw):
            if not keep_piece(kind, piece):
                continue
            for incident in maybe_split_incidents(piece):
                if is_no_issue(incident):
                    continue
                row = base_row(meta)
                row.属地 = cell_value(grid, r_index, mapping.get("属地"), carry=True)
                row.责任区域 = cell_value(grid, r_index, mapping.get("责任区域"), carry=True)
                row.第一责任人 = cell_value(grid, r_index, mapping.get("第一责任人"), carry=True)
                row.责任人 = cell_value(grid, r_index, mapping.get("责任人"), carry=True)
                row.事件描述 = tighten(incident)
                apply_common(
                    row,
                    meta,
                    cell_value(grid, r_index, mapping.get("巡查类型"), carry=True),
                    cell_value(grid, r_index, mapping.get("隐患分级"), carry=True),
                    cell_value(grid, r_index, mapping.get("厂区"), carry=True),
                    cell_value(grid, r_index, mapping.get("日期"), carry=True),
                )
                output.append(row)
    return output


def rows_from_checklist(meta: MailMeta, grid, header_index: int, mapping: dict[str, int]) -> list[Row]:
    output: list[Row] = []
    for r_index in range(header_index + 1, len(grid)):
        serial = cell_value(grid, r_index, mapping.get("序号"), carry=False)
        if serial and (len(tighten(serial)) > 6 or "checklist" in serial.lower() or serial.startswith(("一、", "二、", "三、"))):
            break
        desc_raw = cell_value(grid, r_index, mapping.get("事件描述"), carry=False)
        if desc_raw is None or not tighten(desc_raw):
            continue
        building = tighten(cell_value(grid, r_index, mapping.get("属地"), carry=True))
        place = tighten(cell_value(grid, r_index, mapping.get("巡查位置"), carry=True))
        blocks = split_buildings(desc_raw)
        expanded: list[tuple[str, str]] = []
        for block_name, block_text in blocks:
            area = block_name or (place or building)
            for kind, piece in split_numbered(block_text):
                if not keep_piece(kind, piece):
                    continue
                # 没有编号的长段落按句号拆开，无异常的句子丢掉。
                if kind == "u":
                    sentences = [part.strip() for part in re.split(r"(?<=。)", piece) if part.strip()]
                    if len(sentences) > 1:
                        for sentence in sentences:
                            if keep_piece("u", sentence) and not closed_change(sentence):
                                expanded.append((area, sentence))
                        continue
                if not closed_change(piece):
                    expanded.append((area, piece))
        for area, piece in expanded:
            row = base_row(meta)
            row.属地 = area
            row.责任区域 = ""
            # 清单里的「责任人」是楼层负责人名单，示例汇总表这一列留空。
            row.责任人 = ""
            row.事件描述 = tighten(piece).strip()
            apply_common(row, meta, "", "", "", "")
            if not row.巡查类型:
                row.巡查类型 = "高管驻场"
            output.append(row)
    return output


def dedupe_rows(rows: list[Row]) -> list[Row]:
    seen: set[tuple] = set()
    output: list[Row] = []
    for row in rows:
        key = (
            fingerprint(row.来源),
            fingerprint(row.属地),
            fingerprint(row.责任区域),
            fingerprint(row.事件描述),
        )
        if not key[-1] or key in seen:
            continue
        seen.add(key)
        output.append(row)
    return output


def message_included(meta: MailMeta, start: date | None, end: date | None, overnight: bool = True) -> tuple[bool, str]:
    if start is None or end is None:
        return True, "未限定日期，全部纳入"
    content = meta.content_dates
    hit = [value for value in content if start <= value <= end]
    if hit:
        return True, "文件名/主题/值班日期/表格日期落在范围内: " + ",".join(d.isoformat() for d in hit)
    sent = meta.sent_date
    duty = any(token in (meta.subject + meta.body[:400]) for token in ("值班", "驻厂", "巡查"))
    if overnight and sent and sent == end + timedelta(days=1) and meta.sent_hour is not None and meta.sent_hour < 12 and duty:
        if content and all(value > end + timedelta(days=1) or value < start for value in content):
            return False, "次日上午发出，但正文日期已超出范围次日"
        return True, f"次日 {sent.isoformat()} {meta.sent_hour:02d} 点前收到的值班/驻厂邮件，按夜班补入 {end.isoformat()}"
    shown = ",".join(d.isoformat() for d in content[:6]) or "无"
    sent_text = sent.isoformat() if sent else "无"
    return False, f"不在范围内（识别到的日期: {shown}；发送日: {sent_text}）"


def assign_dates(rows: list[Row], meta: MailMeta, start: date | None, end: date | None) -> None:
    anchor = duty_anchor(meta)
    single = start is not None and start == end
    for row in rows:
        if single:
            row.日期 = start
            continue
        if row.日期 is None:
            row.日期 = anchor


def select_and_parse(
    metas: list[MailMeta],
    start: date | None,
    end: date | None,
    overnight: bool = True,
) -> tuple[list[Row], list[str]]:
    chosen: list[MailMeta] = []
    logs: list[str] = []
    included_bases = set()
    for meta in metas:
        ok, reason = message_included(meta, start, end, overnight)
        if not ok:
            logs.append(f"跳过  {meta.path.name}  → {reason}")
            continue
        if meta.is_reply:
            logs.append(f"待定  {meta.path.name}  → 答复邮件，若原信已纳入则跳过。{reason}")
        else:
            included_bases.add(fingerprint(meta.base_subject))
        chosen.append(meta)
    rows: list[Row] = []
    for meta in chosen:
        if meta.is_reply and fingerprint(meta.base_subject) in included_bases:
            logs.append(f"跳过  {meta.path.name}  → 原信已在范围内，答复只是转发/引用，不重复计数")
            continue
        try:
            parsed = parse_mail_rows(meta)
        except Exception:
            logs.append(f"失败  {meta.path.name}\n{traceback.format_exc()}")
            continue
        assign_dates(parsed, meta, start, end)
        if not any(item.startswith(f"纳入  {meta.path.name}") for item in logs):
            ok, reason = message_included(meta, start, end, overnight)
            sender_note = meta.sender
            if sender_note and not re.search(r"[\u4e00-\u9fff]", sender_note):
                sender_note += "（这封邮件的发件人和签名都没有中文名）"
            logs.append(f"纳入  {meta.path.name}  → {len(parsed)} 条。{reason}。发件人: {sender_note}")
        rows.extend(parsed)
    return rows, logs


def write_workbook(rows: list[Row], output: Path, start: date | None, end: date | None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.styles.colors import Color
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title(start, end)
    header_font = Font(name="微软雅黑", size=9, bold=True, color=Color(theme=1))
    body_font = Font(name="微软雅黑", size=8, color=Color(theme=1))
    header_fill = PatternFill(fill_type="solid", fgColor=Color(theme=8, tint=0.3999755851924192))
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for col, name in enumerate(COLUMNS, start=1):
        cell = sheet.cell(1, col, name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = alignment
        cell.border = thin
    sheet.row_dimensions[1].height = 28.5
    for r_index, row in enumerate(rows, start=2):
        values = row.as_dict()
        for c_index, name in enumerate(COLUMNS, start=1):
            value = values[name]
            cell = sheet.cell(r_index, c_index, None if value in ("", None) else value)
            if name == "日期" and isinstance(value, date):
                cell.value = datetime.combine(value, time())
                cell.number_format = "mm-dd-yy"
            cell.font = body_font
            cell.alignment = alignment
            cell.border = thin
        desc_len = len(row.事件描述 or "")
        lines = max(1, (desc_len + 26) // 27)
        sheet.row_dimensions[r_index].height = max(18.75, min(lines * 15, 90))
    for letter, width in COL_WIDTHS.items():
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{max(1, len(rows) + 1)}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.horizontalCentered = True
    sheet.sheet_view.view = "pageBreakPreview"
    sheet.sheet_view.zoomScale = 70
    sheet.print_title_rows = "1:1"
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.5
    sheet.page_margins.bottom = 0.5
    sheet.page_margins.header = 0.2
    sheet.page_margins.footer = 0.2
    sheet.oddHeader.left.text = f"&8&K7F8C8F值班汇总  {sheet_title(start, end)}"
    sheet.oddFooter.left.text = "&8内部资料"
    sheet.oddFooter.right.text = "&8第 &P 页 / 共 &N 页"
    sheet.auto_filter.ref = None
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def load_check_rows(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    headers = [sheet.cell(1, c).value for c in range(1, 13)]
    rows = []
    for r_index in range(2, sheet.max_row + 1):
        item = {}
        empty = True
        for c_index, name in enumerate(headers, start=1):
            if name not in COLUMNS:
                continue
            value = sheet.cell(r_index, c_index).value
            if isinstance(value, datetime):
                value = value.date()
            if value not in (None, ""):
                empty = False
            item[name] = value
        if not empty:
            rows.append(item)
    return rows


def comparable(value) -> str:
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    return fingerprint(value)


def check_against(rows: list[Row], example_path: Path) -> str:
    example = load_check_rows(example_path)
    used: set[int] = set()
    matched = []
    missed = []
    # 先配长描述，避免短句被大段文本提前吃掉。
    order = sorted(range(len(example)), key=lambda i: len(fingerprint(example[i].get("事件描述"))), reverse=True)
    generated = list(rows)
    for index in order:
        item = example[index]
        target = fingerprint(item.get("事件描述"))
        if not target:
            continue
        best = None
        best_score = 0.0
        for g_index, row in enumerate(generated):
            if g_index in used:
                continue
            got = fingerprint(row.事件描述)
            if not got:
                continue
            if got == target:
                score = 1.0
            elif target in got or got in target:
                score = min(len(target), len(got)) / max(len(target), len(got))
            else:
                score = 0.0
            same_source = fingerprint(item.get("来源")) == fingerprint(row.来源)
            if same_source:
                score += 0.05
            if score > best_score:
                best_score = score
                best = g_index
        if best is not None and best_score >= 0.72:
            used.add(best)
            matched.append((item, generated[best]))
        else:
            missed.append(item)
    fields = ["来源", "发件人", "巡查类型", "日期", "厂区", "属地", "责任区域", "第一责任人", "责任人", "事件描述"]
    lines = [
        f"示例 {len(example)} 行，生成 {len(rows)} 行，对上描述 {len(matched)} 行，示例未对上 {len(missed)} 行。",
    ]
    for name in fields:
        ok = 0
        for item, row in matched:
            if comparable(item.get(name)) == comparable(getattr(row, name)):
                ok += 1
        lines.append(f"  {name}: {ok}/{len(matched) if matched else 0}")
    graded = [(item, row) for item, row in matched if row.隐患分级]
    grade_ok = sum(1 for item, row in graded if comparable(item.get("隐患分级")) == comparable(row.隐患分级))
    lines.append(f"  隐患分级(只统计邮件里写了分级的行): {grade_ok}/{len(graded) if graded else 0}")
    type_ok = sum(1 for item, row in matched if comparable(item.get("隐患类型")) == comparable(row.隐患类型))
    lines.append(f"  隐患类型(简单分类，不要求准): {type_ok}/{len(matched) if matched else 0}")
    if missed:
        lines.append("示例里没有对上的描述（最多 15 条）:")
        for item in missed[:15]:
            lines.append(f"  - [{item.get('来源')}] {tighten(str(item.get('事件描述') or ''))[:80]}")
    extra = len(rows) - len(used)
    lines.append(f"生成结果中多出来的行: {extra}（邮件里有、示例没收，或一格拆成了多条）")
    return "\n".join(lines)


def parse_day(text: str | None) -> date | None:
    if not text:
        return None
    found = parse_dates(text)
    if not found:
        raise argparse.ArgumentTypeError(f"无法识别日期: {text}，请用 2026-07-09")
    return found[0]


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="汇总各厂区值班/巡查 .msg，生成与示例相同版式的 Excel")
    parser.add_argument("--input", required=True, help="放有 .msg 的文件夹")
    parser.add_argument("--output", required=True, help="输出的 xlsx 路径")
    parser.add_argument("--start", type=parse_day, help="开始日期，含当天，例如 2026-07-09")
    parser.add_argument("--end", type=parse_day, help="结束日期，含当天。只查一天时与 --start 相同")
    parser.add_argument("--check", help="可选。人工示例 xlsx，用来打印逐列对照")
    parser.add_argument("--no-overnight", action="store_true", help="不要把次日中午 12 点前的夜班邮件补进来")
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_arg_parser()
    args = parser.parse_args(argv)
    folder = Path(args.input)
    if not folder.is_dir():
        print(f"找不到文件夹: {folder}", file=sys.stderr)
        return 2
    start = args.start
    end = args.end or args.start
    if start and end and end < start:
        print("结束日期不能早于开始日期", file=sys.stderr)
        return 2
    metas = load_messages(folder)
    rows, logs = select_and_parse(metas, start, end, overnight=not args.no_overnight)
    output = Path(args.output)
    write_workbook(rows, output, start, end)
    print(f"扫描 {len(metas)} 封 .msg，写入 {len(rows)} 行 → {output}")
    for line in logs:
        print(line)
    if args.check:
        print(check_against(rows, Path(args.check)))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

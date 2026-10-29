#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""给值班汇总 Excel 的最后一列「隐患类型」做判断。

读入汇总脚本生成的 xlsx（或任何带「事件描述」列的表），只改隐患类型，其他列原样保留。
类别和示例表一致。判断顺序是按这批人工标注整理的：先看具体缺陷，再看泄漏、整洁和设备。

用法:
    python classify_hazard_type.py --input 汇总_2026-07-09.xlsx --output 汇总_2026-07-09_已分类.xlsx
    python classify_hazard_type.py --input 汇总.xlsx --check 值班报告汇总/汇总.xlsx
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import Counter
from pathlib import Path

# 与示例「汇总.xlsx」最后一列相同的类别名。
LABELS = (
    "高风险操作核查",
    "设备设施隐患识别",
    "行为/操作",
    "静电",
    "桶/槽子操作",
    "标识标牌",
    "跑冒滴漏",
    "跑冒滴漏（预防）",
    "气味扩散",
    "设备异常",
    "文件记录",
    "物料状态",
    "现场5S",
    "其他",
    "氮气保护",
    "现场状态",
)

DESC_HEADERS = ("事件描述", "巡查发现", "问题描述", "发现项")
TYPE_HEADERS = ("隐患类型", "隐患分类")


def normalize(text: str | None) -> str:
    """去掉空白和句首编号，方便和表格里的写法对齐。"""
    sample = re.sub(r"\s+", "", str(text or ""))
    sample = sample.replace("\xa0", "")
    return re.sub(r"^[0-9]{1,2}[.、．]", "", sample)


def _has(text: str, pattern: str) -> bool:
    return re.search(pattern, text) is not None


def judge(text: str | None) -> tuple[str, str]:
    """返回 (隐患类型, 一句话理由)。空描述记为其他。"""
    sample = normalize(text)
    if not sample:
        return "其他", "描述是空的"

    if _has(sample, r"刷视频|抽烟|返岗|未佩戴|没有佩戴|不佩戴|操作习惯|操作规范|行为规范|拉链未拉|审批单|复核项|物料暴露|无记录"):
        return "行为/操作", "人员行为或记录未执行"

    if _has(sample, r"无静电导出|缺少静电导出|静电接地卡子串|未连接到墙上|静电扁铁起边") and not _has(sample, r"标签|标识"):
        return "静电", "静电导出或接地本身有缺陷"

    # 气源管没关是设施状态，不要因为句子里有「投料」被当成正在作业。
    if _has(sample, r"气源管未关闭"):
        return "设备设施隐患识别", "氮气气源管未关闭"

    if _has(sample, r"冲釜|checklist|直接排放|缺少卡针|流量计浮子|减压浓缩|卸车|掏罐|密闭操作|危化品|溴素|现场跟踪|氮气保护采用直连|敞口桶|坑洼|隔离控制|照明不足|摔倒"):
        return "高风险操作核查", "正在进行高风险操作，或操作条件不完整"
    if _has(sample, r"掏料|投料") and _has(sample, r"氮气|静电|尾气|劳保|防护|气味"):
        return "高风险操作核查", "投料或掏料时核对氮气、静电、尾气或防护"

    if _has(sample, r"缺少.{0,16}标识|标识损坏|标识不完整|标识不全|无使用上限标识|无检查标签|无编号标签|检查标签|国标标识|管路标签|标签损坏|标签标识|缺少标识") and not _has(sample, r"渗漏|漏水|漏油"):
        return "标识标牌", "标识、标签或标牌缺失、损坏"

    if _has(sample, r"合格签|有效期|过期提醒|仪表清单|检查过期") or (_has(sample, r"过期") and _has(sample, r"静电卡子|变送器|仪表")):
        return "文件记录", "合格签、有效期或检查过期"

    if _has(sample, r"氮气保护未开启|缺少氮气尾气|氮气尾气连接"):
        return "氮气保护", "氮气或尾气保护没有接上"

    if _has(sample, r"气味|异味") and not _has(sample, r"渗漏|漏水|滴落|掏料|投料"):
        return "气味扩散", "有气味，但没有写明泄漏点"

    if _has(sample, r"单阀|未上锁|未进行上锁|没有上锁|缺少盲板|无盲板|阀门未关闭|盲板.{0,8}未紧固|视镜内部") and not _has(sample, r"渗漏|漏水|滴落|滴水"):
        return "跑冒滴漏（预防）", "阀门、盲板或锁具不完整，物料还没漏出来"

    if _has(sample, r"转移入库|物料暂存"):
        return "物料状态", "物料待转移、入库或暂存"
    if _has(sample, r"至塑料桶|氨解液"):
        return "桶/槽子操作", "物料打到桶或槽里"

    if _has(sample, r"状态比较差|控制比较差|状态差") and _has(sample, r"清理|清洁"):
        return "现场5S", "现场状态差，要求是清理"
    if _has(sample, r"现场状态比较差|现场控制比较差|状态比较差|控制比较差|待清洗管路"):
        return "现场状态", "现场状态或控制差，没有具体泄漏点"

    if _has(sample, r"变形|折叠形变|弯折|折断|开口没有紧固|直接开口|安装不完善|疏散|可燃气体|检测器下方|油漆|稀释剂紧邻|消音|堵死|四氟|安装倾斜|开启状态|堵塞门口|影响通行|塑料拍|存在损坏|不易辨认|书写内容|材质一致性|液位报警器未恢复|存在快开|螺丝快开|快开处|插底管缺少压力表"):
        return "设备设施隐患识别", "设备安装、完整性或布置有问题"
    if _has(sample, r"渗漏") and _has(sample, r"紧固"):
        return "设备设施隐患识别", "连接处渗漏，处理方式是紧固"

    if _has(sample, r"漏油|压力表损坏|表损坏|疑似漏|腐蚀痕迹|故障灯|故障|松脱|断裂损坏|无法正常使用|报警功能失效|风幕机|压差|温度有上升|手柄损坏|对讲机|报警状态|常开状态"):
        return "设备异常", "设备故障、损坏或功能失效"

    # 夹套结冰后滴下来的冷凝水，示例不记成泄漏。
    if _has(sample, r"结冰|滴冷凝水") and not _has(sample, r"滴漏|清理|清洁"):
        return "其他", "结冰后的冷凝水，不是物料泄漏"
    if _has(sample, r"缓冲间有滴水"):
        return "现场5S", "缓冲间滴水，要求清理"

    if _has(sample, r"渗漏|漏水|漏气|滴水|滴落|淌水|溢水|撒漏|滴漏|渗液|污水|承接夹套|物料痕迹|打压试密|垫圈"):
        return "跑冒滴漏", "已经有物料或水漏出、滴落或渗出"
    if _has(sample, r"析出") and not _has(sample, r"表面"):
        return "跑冒滴漏", "接口或泵出口有物料析出"
    if _has(sample, r"析出") and _has(sample, r"试密|垫片|地面有"):
        return "跑冒滴漏", "析出并要求试密或清理地面物料"

    if _has(sample, r"清理|清洁|粉尘|锈|摆放较乱|杂乱|未包保温|保温脱落|地漏|抹布|印记|收起|冷凝水|析出|沾染|应放置|滤芯|放置于地面"):
        return "现场5S", "现场不整洁，要求清理、复位或保温"

    if _has(sample, r"无危险操作|无风险|反应平稳|交班|试密已经合格|密闭出料|设备清洗|准备转移|偶联反应|烘料完成|收料作业已经完成|无异常"):
        return "其他", "正常生产状态，或没有单独归类的缺陷"
    if _has(sample, r"开口|缺少液位计|断开|未关闭|缺少专用"):
        return "设备设施隐患识别", "缺件、断开、开口或未关闭"
    return "其他", "没有命中更具体的类别"


def classify_hazard_type(text: str | None) -> str:
    label, _reason = judge(text)
    return label


def _header_map(sheet) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, sheet.max_column + 1):
        value = sheet.cell(1, col).value
        if value is None:
            continue
        mapping[str(value).strip()] = col
    return mapping


def _find_desc_col(headers: dict[str, int]) -> int:
    for name in DESC_HEADERS:
        if name in headers:
            return headers[name]
    raise SystemExit("表头里没有「事件描述」列，无法判断隐患类型。")


def _find_type_col(sheet, headers: dict[str, int]) -> int:
    for name in TYPE_HEADERS:
        if name in headers:
            return headers[name]
    col = sheet.max_column + 1
    header = sheet.cell(1, col, "隐患类型")
    if col > 1:
        source = sheet.cell(1, col - 1)
        header.font = source.font.copy()
        header.alignment = source.alignment.copy()
        header.border = source.border.copy()
        header.fill = source.fill.copy()
    return col


def fill_workbook(source: Path, destination: Path) -> list[tuple[int, str, str, str]]:
    """只改隐患类型列。返回 (行号, 描述, 类型, 理由)。"""
    from openpyxl import load_workbook

    workbook = load_workbook(source)
    sheet = workbook.active
    headers = _header_map(sheet)
    desc_col = _find_desc_col(headers)
    type_col = _find_type_col(sheet, headers)
    rows: list[tuple[int, str, str, str]] = []
    for index in range(2, sheet.max_row + 1):
        desc = sheet.cell(index, desc_col).value
        if desc is None or str(desc).strip() == "":
            continue
        label, reason = judge(str(desc))
        cell = sheet.cell(index, type_col, label)
        if type_col <= len(headers):
            # 原有单元格保留字体和对齐，只换文字。
            pass
        else:
            neighbor = sheet.cell(index, desc_col)
            cell.font = neighbor.font.copy()
            cell.alignment = neighbor.alignment.copy()
            cell.border = neighbor.border.copy()
        rows.append((index, str(desc), label, reason))
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return rows


def _fingerprint(text: str | None) -> str:
    sample = normalize(text).lower()
    return re.sub(r"[，。；、：,.;:（）()\-—]", "", sample)


def load_labeled_rows(path: Path) -> list[tuple[str, str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    headers = _header_map(sheet)
    desc_col = _find_desc_col(headers)
    type_col = None
    for name in TYPE_HEADERS:
        if name in headers:
            type_col = headers[name]
            break
    if type_col is None:
        raise SystemExit(f"{path} 没有隐患类型或隐患分类列，不能当作对照。")
    rows = []
    for index in range(2, sheet.max_row + 1):
        desc = sheet.cell(index, desc_col).value
        label = sheet.cell(index, type_col).value
        if desc is None or str(desc).strip() == "" or label is None or str(label).strip() == "":
            continue
        rows.append((str(desc), str(label).strip()))
    return rows


def score_against(labeled: list[tuple[str, str]]) -> tuple[int, int, list[str]]:
    matched = 0
    misses = []
    for desc, label in labeled:
        predicted = classify_hazard_type(desc)
        if predicted == label:
            matched += 1
        else:
            misses.append(f"[{label} -> {predicted}] {normalize(desc)[:80]}")
    return matched, len(labeled), misses


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="判断值班汇总表最后一列隐患类型")
    parser.add_argument("--input", required=True, help="已生成的汇总 xlsx")
    parser.add_argument("--output", help="写出的 xlsx。不写则在原文件名后加 _已分类")
    parser.add_argument("--check", help="带人工隐患类型的 xlsx，用来打印判对了多少")
    parser.add_argument("--explain", action="store_true", help="在终端打印每条的类型和理由")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    source = Path(args.input)
    if not source.is_file():
        print(f"找不到文件: {source}", file=sys.stderr)
        return 2
    destination = Path(args.output) if args.output else source.with_name(source.stem + "_已分类.xlsx")
    judged = fill_workbook(source, destination)
    counts = Counter(label for _index, _desc, label, _reason in judged)
    print(f"已判断 {len(judged)} 行，写入 {destination}")
    for label in LABELS:
        if counts[label]:
            print(f"  {label}: {counts[label]}")
    other = sum(n for label, n in counts.items() if label not in LABELS)
    if other:
        print(f"  未列出的类别: {other}")
    if args.explain:
        for index, desc, label, reason in judged:
            print(f"R{index} {label}｜{reason}｜{normalize(desc)[:60]}")
    if args.check:
        labeled = load_labeled_rows(Path(args.check))
        matched, total, misses = score_against(labeled)
        print(f"对照 {args.check}：有标注 {total} 行，判对 {matched} 行。")
        for line in misses[:20]:
            print("  " + line)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
